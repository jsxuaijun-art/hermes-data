#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""《盈信财税 · 合规账服务报价表》docx 生成器（compliant-accounting 技能自带）

用法：
    python3 gen_quote.py [--out "输出路径.docx"] [--logo xiangyun.png]

数据来源：D:\\OneDrive\\Desktop\\合规\\01+合规账产品手册_带logo.docx（公司现行手册）
纪律：以下常量全部逐字取自手册原文，改价必须先改手册/参考文件，再改这里的常量。
     手册未载明的项目一律留白，禁止自行定价或估算。
同源文档：references/13-合规账报价表与定价口径.md（改一处必须同步另一处）
"""
import argparse
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

# ══════════════════════════════════════════════════════════════
# 一、内容常量（改文案/改价只动这里）
# ══════════════════════════════════════════════════════════════

COMPANY = "苏州盈信企业管理有限公司"
TITLE = "合规账服务报价表"
SUBTITLE = ("苏州盈信企业管理有限公司    |    国家税务总局 TSC 五级涉税专业服务机构    "
            "|    苏州市财政局备案")
TEL_LINE = "电话：132-2229-7318    180-1262-7126"
SERVICE_LINE = "服务内容：公司注册 · 代理记账 · 合规账 · 税务合规咨询 · 高企账务与申报"

QUALIFICATIONS = [
    "成立于 2009 年，深耕财税服务 17 年，苏州（总部）· 上海（分部）两地服务。",
    "国家税务总局涉税专业服务信用评价 TSC 五级（438.11 分）——涉税服务机构信用评价最高等级。",
    "苏州市财政局备案代理记账机构，逐年参加年检，财政局官网可查。",
    "高级会计师领衔的专业团队，配备中级会计师、税务师，多人具有 8 年以上财税经验。",
    "服务对象：苏州、上海地区中小微企业。",
]

# 合规账产品阶梯（手册原表：营业额 → 月费）
PRODUCTS = [
    ["会计", "0 – 100 万", "1,000 元/月", "每年上门服务 1 次 · 微信 1 小时内响应 · 3 年以上经验会计做账"],
    ["会计", "100 – 300 万", "1,500 元/月", "每年上门服务 1 次 · 微信 1 小时内响应"],
    ["会计", "300 万以上", "建议升级会计主管套餐", "—"],
    ["会计主管", "0 – 300 万", "2,000 元/月", "每季度上门 1 次 · 5 年经验会计做账 + 8 年经验会计师审核"],
    ["会计主管", "300 – 500 万", "3,000 元/月", "每季度上门 1 次 · 微信 1 小时内响应"],
    ["会计主管", "500 – 1,000 万", "4,000 元/月", "每季度上门 1 次 · 微信 1 小时内响应"],
    ["会计主管", "1,000 万以上", "建议升级财务经理套餐", "—"],
    ["财务经理", "1,000 – 2,000 万", "5,000 元/月", "每月上门 1 次 · 注册会计师、税务师审核"],
    ["财务经理", "2,000 – 3,000 万", "7,000 元/月", "每月上门 1 次 · 注册会计师、税务师审核"],
    ["财务经理", "3,000 – 5,000 万", "8,000 元/月", "每月上门 1 次 · 注册会计师、税务师审核"],
    ["财务经理", "5,000 万以上", "协议定价", "按业务复杂度单独约定"],
    ["共享财务总监", "不限（按需）", "15.98 万元/年起", "叠加财务经理全部服务 + 内控、资金、融资、股权、预算"],
]

# 服务内容对照矩阵
MATRIX = [
    ["会计核算（记账、报表、纳税申报）", "√", "√", "√", "√"],
    ["涉税服务（税务申报、风险提示）", "—", "√", "√", "√"],
    ["财务分析（经营数据解读与建议）", "—", "—", "√", "√"],
    ["财务内控（制度、流程、审批）", "—", "—", "—", "√"],
    ["资金管理", "—", "—", "—", "√"],
    ["业财融合", "—", "—", "—", "√"],
    ["融资辅助", "—", "—", "—", "√"],
    ["股权架构", "—", "—", "—", "√"],
    ["预算管理", "—", "—", "—", "√"],
]

# 税务合规产品（可单独选购）
TAX_PRODUCTS = [
    ["税务日常管理", "一年两次税务体检、常年税务咨询、业务流程涉税风险提示", "√", "√", "√"],
    ["合规体系建设", "发票管理规范、税务风险控制、重点风险事项警示", "—", "√", "√"],
    ["税务专项咨询", "税收优惠政策适用评估、税务稽查应对方案、股权架构涉税", "—", "—", "√"],
    ["税务专项服务", "税务稽查应对、税收优惠申请、税务尽职调查（按项目单独报价）", "—", "—", "√（单项报价）"],
]
TAX_FEES = ["税务主管\n6,800 元/年", "税务经理\n9,800 元/年", "税务总监\n12,800 元/年"]

# 开业一次性费用
ONETIME = [
    ["地址费用", "因区域、产权性质、用途不同而异（自有地址不产生）", "1,500 – 4,000 元"],
    ["印章刻制", "公章、法人章、财务专用章等（刻字店收取）", "300 元"],
    ["代办服务费", "签订代理记账 / 合规账服务协议可免", "800 元（可免）"],
    ["银行开户", "代预约协调免费；全程代办", "免费 / 50 元 / 500 元"],
]

QUOTE_NOTES = [
    "上述收费标准为手册规定的建议区间；最终报价依据企业年营业额、票据量、行业复杂度、账务历史遗留问题综合确定。",
    "已签约普通代理记账客户升级合规账：原服务剩余期间按比例折算抵扣。",
    "公司设立登记类费用（地址、刻章、开户）属代收代付，不作为我司服务收入。",
    "多主体、多门店客户可另行商定打包价；共享财务总监档按协议约定。",
    "本表为参考报价，正式报价以双方确认的《合规账服务方案》为准。",
]

# 服务响应承诺
COMMIT = [
    ["上门服务频率", "每年 1 次", "每季度 1 次", "每月 1 次"],
    ["微信响应时效", "1 小时内", "1 小时内", "1 小时内"],
    ["做账人员配置", "3 年以上经验会计", "5 年经验会计", "5 年经验会计"],
    ["审核人员配置", "—", "8 年经验会计师", "注册会计师、税务师"],
]
COMMIT_HEAD = ["服务方式", "会计档", "会计主管档", "财务经理档"]

# ══════════════════════════════════════════════════════════════
# 二、样式
# ══════════════════════════════════════════════════════════════

FONT = "微软雅黑"
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
NAVY_HEX = "1F3A5F"
GRAY = RGBColor(0x59, 0x59, 0x59)
RED = RGBColor(0xC0, 0x00, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ROW_ALT = "F2F5F9"
BAND = "E8EEF5"
TEXT_WIDTH = 17.8  # cm，A4 21cm − 左右边距 3.2cm


def set_run(run, size: float = 10.5, bold: bool = False, color=None, name: str = FONT):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rPr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        rf.set(qn(a), name)


def para(doc, text="", size: float = 10.5, bold: bool = False, color=None, align=None,
         space_before: float = 0, space_after: float = 4, line: float = 1.25, indent=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line
    if align is not None:
        p.alignment = align
    if indent is not None:
        pf.left_indent = Cm(indent)
    if text:
        set_run(p.add_run(text), size=size, bold=bold, color=color)
    return p


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def cell_margins(cell, top=40, bottom=40, left=80, right=80):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for tag, val in (("top", top), ("start", left), ("bottom", bottom), ("end", right)):
        e = OxmlElement(f"w:{tag}")
        e.set(qn("w:w"), str(val))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    tcPr.append(mar)


def vcenter(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    va = OxmlElement("w:vAlign")
    va.set(qn("w:val"), "center")
    tcPr.append(va)


def fill_cell(cell, text, size: float = 10, bold: bool = False, color=None,
              align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(1.5)
    p.paragraph_format.space_after = Pt(1.5)
    p.paragraph_format.line_spacing = 1.15
    # 支持 \n 多行（表头换行）
    for i, seg in enumerate(str(text).split("\n")):
        if i:
            p.add_run().add_break()
        set_run(p.add_run(seg), size=size, bold=bold, color=color)
    vcenter(cell)


def make_table(doc, headers, rows, widths, price_col=None, header_size: float = 9.5,
               body_size: float = 9.5):
    assert abs(sum(widths) - TEXT_WIDTH) < 0.05, f"列宽合计 {sum(widths)} ≠ 版心 {TEXT_WIDTH}cm"
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, (h, w) in enumerate(zip(headers, widths)):
        c = t.rows[0].cells[i]
        c.width = Cm(w)
        fill_cell(c, h, size=header_size, bold=True, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
        shade(c, NAVY_HEX)
        cell_margins(c)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for ci, val in enumerate(row):
            c = cells[ci]
            c.width = Cm(widths[ci])
            is_price = price_col is not None and ci == price_col
            fill_cell(c, val, size=body_size, bold=is_price or ci == 0,
                      color=RED if is_price else None,
                      align=WD_ALIGN_PARAGRAPH.CENTER if (is_price or ci != len(row) - 1)
                      else WD_ALIGN_PARAGRAPH.LEFT)
            cell_margins(c)
            if ri % 2 == 1:
                shade(c, ROW_ALT)
    return t


def section_title(doc, text, space_before=12):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(5)
    pf.line_spacing = 1.1
    pf.left_indent = Cm(0.05)
    set_run(p.add_run("▍"), size=12.5, bold=True, color=RED)
    set_run(p.add_run(text), size=12.5, bold=True, color=NAVY)
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bt = OxmlElement("w:bottom")
    bt.set(qn("w:val"), "single")
    bt.set(qn("w:sz"), "6")
    bt.set(qn("w:space"), "3")
    bt.set(qn("w:color"), NAVY_HEX)
    pbdr.append(bt)
    pPr.append(pbdr)
    return p


def bullets(doc, items, size=9.5, indent=0.35):
    for t in items:
        para(doc, "· " + t, size=size, space_after=2.5, indent=indent)


# ══════════════════════════════════════════════════════════════
# 三、装配
# ══════════════════════════════════════════════════════════════

def build(out: Path, logo: Path | None = None):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(1.6), Cm(1.4)
    sec.left_margin = sec.right_margin = Cm(1.6)

    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    hdr = sec.header.paragraphs[0]
    hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run(hdr.add_run(f"{COMPANY}  ·  合规账服务"), size=8.5, color=GRAY)
    if logo and Path(logo).exists():
        hp = sec.header.add_paragraph()
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hp.add_run().add_picture(str(logo), width=Cm(0.8))

    para(doc, TITLE, size=21, bold=True, color=NAVY,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2, line=1.1)
    para(doc, SUBTITLE, size=9, color=GRAY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)

    section_title(doc, "一、客户信息（请填写）", space_before=0)
    info = doc.add_table(rows=2, cols=4)
    info.style = "Table Grid"
    info.autofit = False
    labels = ["客户名称", "联系人 / 电话", "所属行业", "报价日期"]
    w = [5.0, 4.6, 4.6, 3.6]
    for i, (lb, ww) in enumerate(zip(labels, w)):
        c = info.rows[0].cells[i]
        c.width = Cm(ww)
        fill_cell(c, lb, size=9.5, bold=True, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER)
        shade(c, BAND)
        cell_margins(c)
    for i, ww in enumerate(w):
        c = info.rows[1].cells[i]
        c.width = Cm(ww)
        fill_cell(c, "", size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        cell_margins(c)

    section_title(doc, "二、服务方资质")
    bullets(doc, QUALIFICATIONS, size=10)

    section_title(doc, "三、合规账产品与收费标准（核心）")
    para(doc, "合规账以「年营业额」为主定价维度，四档产品逐级叠加：会计 → 会计主管 → 财务经理 → 共享财务总监。",
         size=9.5, color=GRAY, space_after=5)
    make_table(doc, ["产品档位", "适用年营业额", "建议收费标准", "服务方式与人员配置"],
               PRODUCTS, [2.3, 2.9, 3.5, 9.1], price_col=2)
    para(doc, "注：表中为「建议收费标准」区间（手册原表），最终报价按票据量、行业复杂度、账务历史遗留问题等综合确定。",
         size=8.5, color=GRAY, space_before=4, space_after=2)

    section_title(doc, "四、各档服务内容对照")
    make_table(doc, ["服务模块", "会计", "会计主管", "财务经理", "共享财务总监"],
               MATRIX, [6.4, 2.8, 2.8, 2.9, 2.9])

    section_title(doc, "五、税务合规产品（可单独选购，与上述套餐叠加）")
    make_table(doc, ["服务类别", "服务内容"] + TAX_FEES, TAX_PRODUCTS,
               [2.9, 7.0, 2.55, 2.55, 2.8], body_size=9.0, header_size=9.0)
    para(doc, "√＝包含　—＝不含。三档含金量逐级提升：主管档打基础，经理档建体系，总监档覆盖专项与稽查应对。",
         size=8.5, color=GRAY, space_before=4, space_after=2)

    section_title(doc, "六、开业一次性费用（公司设立登记，代收代付项目）")
    make_table(doc, ["项目", "说明", "参考金额"], ONETIME, [3.2, 9.0, 5.6], price_col=2)

    section_title(doc, "七、标准代理记账（对照参考）")
    para(doc, "一般纳税人：500 元/月起；零申报：100 元/月。合规账在标准记账之上增加风险排查、合规建议与上门服务，"
              "两者价差对应的正是「合规」部分的专业投入。", size=10, space_after=2)

    section_title(doc, "八、报价说明")
    bullets(doc, QUOTE_NOTES)

    section_title(doc, "九、服务响应承诺")
    make_table(doc, COMMIT_HEAD, COMMIT, [3.6, 4.6, 4.6, 5.0])

    section_title(doc, "十、联系方式")
    para(doc, f"{COMPANY}（苏州总部 · 上海分部）", size=10.5, bold=True, color=NAVY, space_after=2)
    para(doc, TEL_LINE, size=10.5, space_after=2)
    para(doc, SERVICE_LINE, size=9.5, color=GRAY, space_after=2)
    para(doc, "（本报价表依据《合规账产品手册》编制；手册未载明事项，已在表中如实留白，未作估计。）",
         size=8.5, color=GRAY, space_before=6)

    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return out


def main():
    here = Path(__file__).resolve().parent
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=r"C:\Users\Administrator\Desktop\盈信财税_合规账服务报价表.docx")
    ap.add_argument("--logo", default=str(here / "xiangyun.png"))
    a = ap.parse_args()
    out = build(Path(a.out), Path(a.logo))
    d = Document(str(out))
    print(f"SAVED: {out}  {out.stat().st_size} bytes")
    print(f"段落 {len(d.paragraphs)} · 表格 {len(d.tables)}")
    for i, t in enumerate(d.tables):
        print(f"  表{i+1}: {len(t.rows)}行 x {len(t.columns)}列")


if __name__ == "__main__":
    main()
